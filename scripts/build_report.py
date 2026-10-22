#!/usr/bin/env python3
"""
build_report.py - Step 7 of the ai-queries-finder skill.

Builds the Excel working sheet from:
  ai-queries.json  (find_ai_queries.py)
  strategy.json    (written by Claude: per-cluster optimisation plan, quick wins, 7/30/90-day plan)

Tabs: Summary | AI Queries | AI Feature Pages | Regex Library | Topic Clusters | Opportunities |
      Optimization Plan | AI Visibility Plan | Method & Notes

strategy.json shape (every key optional):
{
  "site_name": "Acme Flooring",
  "headline": "2-3 sentence summary",
  "clusters": {
    "T1": {"target_url": "...", "page_action": "Update existing page", "diagnosis": "...",
           "answer_block": "40-60 words", "question_headings": ["..."], "faq": [{"q": "...", "a": "..."}],
           "format_block": "...", "schema": ["FAQPage"], "internal_links": ["..."], "off_site": ["..."],
           "notes": "facts to confirm"}
  },
  "quick_wins": [{"action": "...", "why": "...", "effort": "Low", "impact": "High", "timeframe": "Days 1-7"}],
  "plan": [{"phase": "Days 8-30", "action": "...", "kpi": "..."}],
  "measurement": ["..."],
  "pages": {"<url>": "what to do with this page for AI features"}
}

Usage:
  python build_report.py --data work/ai-queries.json --strategy work/strategy.json -o outputs/ai-queries.xlsx
         [--md outputs/summary.md] [--top-opportunities 50]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, INK, MUTED = "0D1730", "1F2937", "6B7280"
CYAN, VIOLET, GREEN, AMBER, RED = "22D3EE", "A78BFA", "34D399", "FBBF24", "F87171"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
SUGGEST_FILL = PatternFill("solid", fgColor="ECFDF5")   # Claude's suggestions = light green
INPUT_FILL = PatternFill("solid", fgColor="FFF7E6")     # columns for the user to fill = light amber
KPI_FILL = PatternFill("solid", fgColor="F1F5F9")
THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
PRIORITY_FILL = {"High": "FEE2E2", "Medium": "FEF3C7", "Low": "E0F2FE"}
TIER_FILL = {"High": "EDE9FE", "Medium": "E0F2FE", "Low": "F3F4F6"}
STATUSES = '"To do,In progress,Done,Skipped"'
NOT_WRITTEN = "(not written yet)"
ATTRIBUTION = "Built with ai-queries-finder by Akif Qureshi | SEO Specialist - algoblueprints.com"
Q = "'AI Queries'"


def table(ws, headers: list[tuple[str, int]], rows: list[list], start_row: int = 1,
          suggest_cols: tuple = (), input_cols: tuple = (), status_col: int | None = None) -> int:
    """Write a styled table. headers = [(name, width)]. Returns last row number."""
    for ci, (name, width) in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=ci, value=name)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[start_row].height = 32
    for ri, row in enumerate(rows, start_row + 1):
        for ci, val in enumerate(row, 1):
            if isinstance(val, Formula):
                c = ws.cell(row=ri, column=ci, value=str(val))
            else:
                c = ws.cell(row=ri, column=ci, value=val)
                if isinstance(val, str) and val.startswith("="):
                    c.data_type = "s"  # text that happens to start with '=' must not become a formula
            c.alignment, c.border = WRAP, BORDER
            c.font = Font(size=10, color=INK)
            if ci in suggest_cols:
                c.fill = SUGGEST_FILL
            elif ci in input_cols:
                c.fill = INPUT_FILL
    last = start_row + len(rows)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{last}"
    if status_col and rows:
        dv = DataValidation(type="list", formula1=STATUSES, allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(status_col)
        dv.add(f"{col}{start_row + 1}:{col}{max(last, start_row + 1)}")
        ws.conditional_formatting.add(f"{col}{start_row + 1}:{col}{last}",
                                      CellIsRule(operator="equal", formula=['"Done"'],
                                                 fill=PatternFill("solid", fgColor="D1FAE5")))
    return last


class Formula(str):
    """Marks a cell value as a real formula (plain strings starting with '=' are kept as text)."""


def fmt(ws, col: str, first: int, last: int, number_format: str) -> None:
    for r in range(first, last + 1):
        ws[f"{col}{r}"].number_format = number_format


def fill_by_value(ws, col: str, first: int, last: int, fills: dict) -> None:
    for r in range(first, last + 1):
        v = ws[f"{col}{r}"].value
        if v in fills:
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=fills[v])
            ws[f"{col}{r}"].font = Font(size=10, bold=True, color=INK)


def bullets(items) -> str:
    if not items:
        return ""
    if isinstance(items, str):
        return items
    return "\n".join(f"- {t}" for t in items)


def faq_text(faq) -> str:
    if not faq:
        return ""
    if isinstance(faq, str):
        return faq
    return "\n\n".join(f"Q: {f.get('q', '')}\nA: {f.get('a', '')}" for f in faq)


def words_formula(cell: str) -> Formula:
    return Formula(f'=IF(LEN(TRIM({cell}))=0,0,LEN(TRIM({cell}))-LEN(SUBSTITUTE(TRIM({cell})," ",""))+1)')


def title_block(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color=MUTED)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", type=Path, required=True)
    ap.add_argument("--strategy", type=Path)
    ap.add_argument("-o", "--out", type=Path, required=True)
    ap.add_argument("--md", type=Path, help="Also write a short markdown summary")
    ap.add_argument("--site-name")
    ap.add_argument("--top-opportunities", type=int, default=50)
    args = ap.parse_args()

    d = json.loads(args.data.read_text(encoding="utf-8"))
    st = json.loads(args.strategy.read_text(encoding="utf-8")) if args.strategy and args.strategy.exists() else {}
    cplan = st.get("clusters", {})
    site_name = args.site_name or st.get("site_name") or urlparse(d["settings"].get("site") or "").netloc or "Website"
    period = d.get("date_range") or "date range not stated (add --date-range to find_ai_queries.py)"
    today = dt.date.today().isoformat()
    tot = d["totals"]
    clusters = d["clusters"]
    queries = d["queries"]

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_q = wb.create_sheet("AI Queries")
    ws_ap = wb.create_sheet("AI Feature Pages")
    ws_rx = wb.create_sheet("Regex Library")
    ws_cl = wb.create_sheet("Topic Clusters")
    ws_op = wb.create_sheet("Opportunities")
    ws_pl = wb.create_sheet("Optimization Plan")
    ws_vp = wb.create_sheet("AI Visibility Plan")
    ws_me = wb.create_sheet("Method & Notes")

    # ---------------------------------------------------------------- AI Queries
    headers = [("Query", 48), ("Words", 7), ("AI-style", 9), ("AI score", 8), ("Tier", 9), ("Regex families", 16),
               ("Intent", 13), ("Topic", 8), ("Clicks", 8), ("Impressions", 11), ("CTR", 8), ("Position", 9),
               ("Bucket", 9), ("AI Overview likely", 11), ("Action type", 30), ("Target CTR", 9),
               ("Est. extra clicks", 10), ("Page", 40)]
    rows = []
    for i, r in enumerate(queries, 2):
        rows.append([r["query"], words_formula(f"A{i}"), "Yes" if r["ai_style"] else "No", r["score"], r["tier"],
                     ", ".join(r["families"]), r["intent"], r["cluster"], r["clicks"], r["impressions"],
                     Formula(f"=IF(J{i}>0,I{i}/J{i},0)"), r["position"], r["bucket"],
                     "Yes" if r["aio_likely"] else "", r["action_type"], r["target_ctr"],
                     Formula(f'=IF(H{i}="",0,MAX(0,ROUND(J{i}*P{i}-I{i},0)))'), r["page"]])
    last_q = table(ws_q, headers, rows)
    fmt(ws_q, "K", 2, last_q, "0.0%")
    fmt(ws_q, "P", 2, last_q, "0.0%")
    fmt(ws_q, "L", 2, last_q, "0.0")
    fmt(ws_q, "J", 2, last_q, "#,##0")
    fill_by_value(ws_q, "E", 2, last_q, TIER_FILL)
    if last_q >= 2:
        ws_q.conditional_formatting.add(f"D2:D{last_q}", ColorScaleRule(start_type="num", start_value=0, start_color="F3F4F6",
                                                                         end_type="num", end_value=100, end_color="A78BFA"))
        ws_q.conditional_formatting.add(f"N2:N{last_q}", CellIsRule(operator="equal", formula=['"Yes"'],
                                                                    fill=PatternFill("solid", fgColor="FEE2E2")))

    def qcol(col: str) -> str:
        return f"{Q}!${col}$2:${col}${max(last_q, 2)}"

    # ---------------------------------------------------------------- AI Feature Pages
    ap_rows = []
    pages = d.get("ai_pages", [])
    paction = st.get("pages", {})
    last_guess = len(pages) + 1
    for i, pg in enumerate(pages, 2):
        ap_rows.append([pg["page"], pg["topic"], pg["ai_impressions"],
                        Formula(f"=IF(SUM($C$2:$C${last_guess})>0,C{i}/SUM($C$2:$C${last_guess}),0)"),
                        pg["impressions"], pg["clicks"], Formula(f'=IF(N(E{i})>0,F{i}/E{i},"")'), pg["position"],
                        Formula(f'=IF(N(E{i})>0,C{i}/E{i},"")'), paction.get(pg["page"], "")])
    if not ap_rows:
        ap_rows.append(["(add --ai-pages to find_ai_queries.py: GSC > Performance > Generative AI features > Pages)",
                        "", 0, 0, "", "", "", "", "", ""])
    last_ap = table(ws_ap, [("Page", 60), ("Topic", 8), ("AI feature impressions", 12), ("Share of AI impressions", 11),
                            ("Search impressions", 12), ("Clicks", 8), ("CTR", 8), ("Avg position", 9),
                            ("AI share of page impressions", 12), ("What to do", 60)], ap_rows, suggest_cols=(10,))
    for col, nf in (("D", "0.0%"), ("G", "0.0%"), ("I", "0.0%"), ("C", "#,##0"), ("E", "#,##0")):
        fmt(ws_ap, col, 2, last_ap, nf)
    if last_ap >= 2:
        ws_ap.conditional_formatting.add(f"C2:C{last_ap}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                                           end_type="max", end_color="A78BFA"))

    # ---------------------------------------------------------------- Regex Library
    title_block(ws_rx, "Regex library", "Paste into GSC > Performance > + Add filter > Query > Custom (regex). "
                                        "RE2 syntax, max 4,096 characters.")
    rx_rows = []
    start = 4
    for i, f in enumerate(d["families"], start + 1):
        rx_rows.append([f["id"], f["name"], f["regex"], Formula(f"=LEN(C{i})"), "Yes" if f["in_master"] else "No",
                        f["weight"], f["catches"],
                        Formula(f'=COUNTIFS({qcol("F")},"*"&A{i}&"*",{qcol("C")},"Yes")'),
                        Formula(f'=SUMIFS({qcol("J")},{qcol("F")},"*"&A{i}&"*",{qcol("C")},"Yes")'),
                        Formula(f'=SUMIFS({qcol("I")},{qcol("F")},"*"&A{i}&"*",{qcol("C")},"Yes")')])
    mi = start + len(rx_rows) + 1
    rx_rows.append(["MASTER", "Master regex (paste this first)", d["master_regex"], Formula(f"=LEN(C{mi})"), "-", "-",
                    "OR of every family marked 'in master'. Use it for the first pull; run single families to get "
                    "past the 1,000-row cap.", Formula(f'=COUNTIF({qcol("C")},"Yes")'),
                    Formula(f'=SUMIFS({qcol("J")},{qcol("C")},"Yes")'),
                    Formula(f'=SUMIFS({qcol("I")},{qcol("C")},"Yes")')])
    last_rx = table(ws_rx, [("ID", 8), ("Family", 26), ("Regex (copy into GSC)", 70), ("Chars", 7), ("In master", 9),
                            ("Weight", 8), ("What it catches", 50), ("AI queries", 10), ("Impressions", 12),
                            ("Clicks", 9)], rx_rows, start_row=start)
    for r in range(start + 1, last_rx + 1):
        ws_rx[f"C{r}"].font = Font(name="Consolas", size=9, color=INK)
    for c in "ABCDEFGHIJ":
        ws_rx[f"{c}{last_rx}"].font = Font(bold=True, size=10, color=NAVY,
                                            name="Consolas" if c == "C" else None)
    fmt(ws_rx, "I", start + 1, last_rx, "#,##0")
    ws_rx.conditional_formatting.add(f"D{start + 1}:D{last_rx}",
                                     CellIsRule(operator="greaterThan", formula=["4096"],
                                                fill=PatternFill("solid", fgColor="FEE2E2")))

    # ---------------------------------------------------------------- Topic Clusters
    cl_rows = []
    for i, c in enumerate(clusters, 2):
        main_intent = next(iter(c["intents"]), "")
        cl_rows.append([c["id"], c["label"], c["priority"],
                        Formula(f'=COUNTIFS({qcol("H")},A{i},{qcol("C")},"Yes")'),
                        Formula(f'=SUMIFS({qcol("I")},{qcol("H")},A{i})'),
                        Formula(f'=SUMIFS({qcol("J")},{qcol("H")},A{i})'),
                        Formula(f"=IF(F{i}>0,E{i}/F{i},0)"), c["position"], c["avg_score"],
                        Formula(f'=COUNTIFS({qcol("H")},A{i},{qcol("N")},"Yes")'),
                        Formula(f'=SUMIFS({qcol("Q")},{qcol("H")},A{i})'),
                        main_intent, ", ".join(f"{k} {v}" for k, v in c["buckets"].items()),
                        "\n".join(c["pages"]), "\n".join(c["top_queries"])])
    last_cl = table(ws_cl, [("ID", 6), ("Topic", 28), ("Priority", 9), ("Queries", 8), ("Clicks", 8),
                            ("Impressions", 11), ("CTR", 8), ("Avg position", 9), ("Avg AI score", 9),
                            ("AI Overview likely", 10), ("Est. extra clicks", 10), ("Main intent", 13),
                            ("Position buckets", 20), ("Pages", 38), ("Top queries", 60)], cl_rows)
    fmt(ws_cl, "G", 2, last_cl, "0.0%")
    fmt(ws_cl, "F", 2, last_cl, "#,##0")
    fill_by_value(ws_cl, "C", 2, last_cl, PRIORITY_FILL)
    if last_cl >= 2:
        ws_cl.conditional_formatting.add(f"F2:F{last_cl}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                                           end_type="max", end_color="22D3EE"))

    # ---------------------------------------------------------------- Opportunities
    opp = sorted((r for r in queries if r["ai_style"] and r["cluster"] and not r.get("agent")),
                 key=lambda r: -max(0.0, r["impressions"] * r["target_ctr"] - r["clicks"]))[:args.top_opportunities]
    op_rows = []
    for i, r in enumerate(opp, 2):
        op_rows.append([r["query"], r["cluster"], r["intent"], r["impressions"], r["clicks"],
                        Formula(f"=IF(D{i}>0,E{i}/D{i},0)"), r["position"], r["bucket"],
                        "Yes" if r["aio_likely"] else "", r["action_type"], r["target_ctr"],
                        Formula(f"=MAX(0,ROUND(D{i}*K{i}-E{i},0))"),
                        Formula(f'=IF(L{i}>=20,"High",IF(L{i}>=5,"Medium","Low"))'), r["page"]])
    last_op = table(ws_op, [("Query", 48), ("Topic", 7), ("Intent", 13), ("Impressions", 11), ("Clicks", 8),
                            ("CTR", 8), ("Position", 9), ("Bucket", 9), ("AI Overview likely", 10), ("Action type", 32),
                            ("Target CTR", 9), ("Est. extra clicks", 10), ("Priority", 9), ("Page", 40)], op_rows)
    fmt(ws_op, "F", 2, last_op, "0.0%")
    fmt(ws_op, "K", 2, last_op, "0.0%")
    fmt(ws_op, "D", 2, last_op, "#,##0")
    for prio, color in PRIORITY_FILL.items():
        if last_op >= 2:
            ws_op.conditional_formatting.add(f"M2:M{last_op}", CellIsRule(operator="equal", formula=[f'"{prio}"'],
                                                                          fill=PatternFill("solid", fgColor=color)))

    # ---------------------------------------------------------------- Optimization Plan
    pl_rows = []
    unwritten = 0
    for i, c in enumerate(clusters, 2):
        p = cplan.get(c["id"], {})
        if not p:
            unwritten += c["priority"] in ("High", "Medium")
        pl_rows.append([c["id"], c["label"], c["priority"], p.get("target_url", c["pages"][0] if c["pages"] else ""),
                        p.get("page_action", ""), p.get("diagnosis", ""), p.get("answer_block", NOT_WRITTEN if
                                                                                  c["priority"] in ("High", "Medium") else ""),
                        words_formula(f"G{i}"), bullets(p.get("question_headings")), faq_text(p.get("faq")),
                        p.get("format_block", ""), ", ".join(p.get("schema", [])) if isinstance(p.get("schema"), list)
                        else p.get("schema", ""), bullets(p.get("internal_links")), bullets(p.get("off_site")),
                        p.get("notes", ""), "To do"])
    last_pl = table(ws_pl, [("Topic", 6), ("Topic name", 22), ("Priority", 9), ("Target URL", 34), ("Page action", 16),
                            ("Diagnosis", 44), ("Answer-first block (40-60 words)", 60), ("Words", 7),
                            ("Question headings", 40), ("FAQ", 60), ("Table / steps / list", 50), ("Schema", 16),
                            ("Internal links", 40), ("Off-site actions", 44), ("Check before publishing", 30),
                            ("Status", 11)], pl_rows,
                    suggest_cols=(6, 7, 9, 10, 11, 12, 13, 14), input_cols=(15,), status_col=16)
    fill_by_value(ws_pl, "C", 2, last_pl, PRIORITY_FILL)
    if last_pl >= 2:
        ws_pl.conditional_formatting.add(f"H2:H{last_pl}", FormulaRule(
            formula=["AND(H2>0,OR(H2<40,H2>60))"], fill=PatternFill("solid", fgColor="FEE2E2")))
    for r in range(2, last_pl + 1):
        ws_pl.row_dimensions[r].height = 150

    # ---------------------------------------------------------------- AI Visibility Plan
    title_block(ws_vp, "AI visibility plan", "Fastest route first: unblock, fix the biggest leaks, then build "
                                              "and compound. Amber = yours to fill in.")
    vp_rows = []
    for w in st.get("quick_wins", []):
        vp_rows.append([w.get("timeframe", "Days 1-7"), "Quick win", w.get("action", ""), w.get("why", ""),
                        w.get("effort", ""), w.get("impact", ""), "", "To do", ""])
    for p in st.get("plan", []):
        vp_rows.append([p.get("phase", ""), "Plan", p.get("action", ""), p.get("kpi", ""), p.get("effort", ""),
                        p.get("impact", ""), "", "To do", ""])
    if not vp_rows:
        vp_rows.append(["Days 1-7", "Plan", NOT_WRITTEN, "", "", "", "", "To do", ""])
    last_vp = table(ws_vp, [("When", 13), ("Type", 10), ("Action", 60), ("Why / KPI", 50), ("Effort", 9),
                            ("Impact", 9), ("Owner", 14), ("Status", 11), ("Date done", 12)], vp_rows, start_row=4,
                    suggest_cols=(3, 4), input_cols=(7, 9), status_col=8)
    for col in ("E", "F"):
        for prio, color in PRIORITY_FILL.items():
            ws_vp.conditional_formatting.add(f"{col}5:{col}{last_vp}", CellIsRule(
                operator="equal", formula=[f'"{prio}"'], fill=PatternFill("solid", fgColor=color)))
    r = last_vp + 2
    ws_vp[f"A{r}"] = "How to measure it"
    ws_vp[f"A{r}"].font = Font(bold=True, size=12, color=NAVY)
    for m in st.get("measurement", []) or ["Re-run the master regex in GSC monthly and compare with this baseline."]:
        r += 1
        ws_vp[f"A{r}"] = "-"
        ws_vp[f"C{r}"] = m
        ws_vp[f"C{r}"].alignment = WRAP
        if isinstance(m, str) and m.startswith("="):
            ws_vp[f"C{r}"].data_type = "s"

    # ---------------------------------------------------------------- Summary
    s = ws_sum
    title_block(s, f"AI Queries Finder - {site_name}", f"Google Search Console, {period}. Built {today}.")
    s.column_dimensions["A"].width = 44
    s.column_dimensions["B"].width = 18
    for c in "CDEFG":
        s.column_dimensions[c].width = 18
    site_clicks = tot.get("site_clicks")
    site_impr = tot.get("site_impressions")
    kpis = [
        ("Site total clicks (no filter)", site_clicks if site_clicks is not None else "not given", "#,##0"),
        ("Site total impressions (no filter)", site_impr if site_impr is not None else "not given", "#,##0"),
        ("AI-style queries found", Formula(f'=COUNTIF({qcol("C")},"Yes")'), "#,##0"),
        ("AI-style impressions", Formula(f'=SUMIFS({qcol("J")},{qcol("C")},"Yes")'), "#,##0"),
        ("AI-style clicks", Formula(f'=SUMIFS({qcol("I")},{qcol("C")},"Yes")'), "#,##0"),
        ("AI-style CTR", Formula("=IF(B7>0,B8/B7,0)"), "0.00%"),
        ("Share of site impressions", Formula('=IF(N(B5)>0,B7/B5,"n/a")'), "0.0%"),
        ("Queries where an AI answer likely takes the click", Formula(f'=COUNTIFS({qcol("C")},"Yes",{qcol("N")},"Yes")'), "#,##0"),
        ("Topic clusters", Formula(f"=COUNTA('Topic Clusters'!$A$2:$A${max(last_cl, 2)})"), "0"),
        ("Est. extra clicks if the plan works (per period)", Formula(f'=SUMIFS({qcol("Q")},{qcol("C")},"Yes")'), "#,##0"),
        ("Impressions in Google AI features (GSC Generative AI report total)",
         tot.get("ai_feature_total") or (Formula(f"=SUM('AI Feature Pages'!$C$2:$C${max(last_ap, 2)})")
                                         if pages else "not pulled"), "#,##0"),
        ("AI Mode follow-up queries (no topic word)", tot.get("followup_queries", 0), "#,##0"),
        ("AI agent / tool searches (site:, intitle:...)", tot.get("agent_queries", 0), "#,##0"),
    ]
    for i, (label, val, nf) in enumerate(kpis, 4):
        s[f"A{i}"], s[f"B{i}"] = label, str(val) if isinstance(val, Formula) else val
        s[f"A{i}"].fill = s[f"B{i}"].fill = KPI_FILL
        s[f"A{i}"].font = Font(size=10, color=INK)
        s[f"B{i}"].font = Font(size=11, bold=True, color=NAVY)
        s[f"B{i}"].number_format = nf
        s[f"A{i}"].border = s[f"B{i}"].border = BORDER
    r = 4 + len(kpis) + 1
    if st.get("headline"):
        s[f"A{r}"] = "The short version"
        s[f"A{r}"].font = Font(bold=True, size=12, color=NAVY)
        s[f"A{r + 1}"] = st["headline"]
        s.merge_cells(f"A{r + 1}:G{r + 1}")
        s[f"A{r + 1}"].alignment = WRAP
        s.row_dimensions[r + 1].height = 64
        r += 3
    s[f"A{r}"] = "Top topics"
    s[f"A{r}"].font = Font(bold=True, size=12, color=NAVY)
    top = [(i, c) for i, c in enumerate(clusters, 2) if c["head"] != "(misc)"][:5]
    trows = [[Formula(f"='Topic Clusters'!B{i}"), Formula(f"='Topic Clusters'!C{i}"),
              Formula(f"='Topic Clusters'!F{i}"), Formula(f"='Topic Clusters'!E{i}"),
              Formula(f"='Topic Clusters'!H{i}"), Formula(f"='Topic Clusters'!K{i}"),
              cplan.get(c["id"], {}).get("page_action", "")] for i, c in top]
    last_t = table(s, [("Topic", 44), ("Priority", 18), ("Impressions", 18), ("Clicks", 18), ("Avg position", 18),
                       ("Est. extra clicks", 18), ("Move", 18)], trows, start_row=r + 1)
    s.freeze_panes = None
    s.auto_filter.ref = None
    fmt(s, "C", r + 2, last_t, "#,##0")
    r = last_t + 2
    wins = st.get("quick_wins", [])[:5]
    if wins:
        s[f"A{r}"] = "Quick wins (days 1-7)"
        s[f"A{r}"].font = Font(bold=True, size=12, color=NAVY)
        for w in wins:
            r += 1
            s[f"A{r}"] = f"- {w.get('action', '')}"
            s.merge_cells(f"A{r}:G{r}")
            s[f"A{r}"].alignment = WRAP
        r += 2
    s[f"A{r}"] = "How to use this sheet"
    s[f"A{r}"].font = Font(bold=True, size=12, color=NAVY)
    for line in ["1. Topic Clusters: what AI-style searchers want from you, biggest first.",
                 "2. Optimization Plan: the answer block, headings, FAQ and links to add per topic (green = review).",
                 "3. AI Visibility Plan: do the quick wins this week, then the 30/90-day plan. Fill the amber cells.",
                 "4. Regex Library: save the master regex as a GSC filter and re-check monthly against this baseline.",
                 "Green cells are suggestions to review. Amber cells are yours to fill in."]:
        r += 1
        s[f"A{r}"] = line
        s.merge_cells(f"A{r}:G{r}")
    s[f"A{r + 2}"] = ATTRIBUTION
    s[f"A{r + 2}"].font = Font(italic=True, size=9, color=MUTED)

    # ---------------------------------------------------------------- Method & Notes
    title_block(ws_me, "Method & notes", ATTRIBUTION)
    ws_me.column_dimensions["A"].width = 34
    ws_me.column_dimensions["B"].width = 100
    notes = [
        ("Date range", period),
        ("Search type", "Web"),
        ("Rows read", tot["rows_read"]),
        ("AI-style threshold", f"AI score >= {d['settings']['min_score']} (tiers: High >= 50, Medium >= 30)"),
        ("Minimum impressions", d["settings"]["min_impr"]),
        ("Brand terms excluded from topics", ", ".join(d["settings"]["brand"]) or "none given"),
        ("Locations (local intent)", ", ".join(d["settings"]["locations"]) or "none given"),
        ("AI score", "Sum of the weights of every regex family the query matches, capped at 100 (see Regex Library)."),
        ("Topics", "Each AI-style, non-brand query joins the topic of its most common meaningful word across all "
                   "AI-style queries. Words used by only one query go to 'Other'."),
        ("Position buckets", "Top 3 / 4-10 / 11-20 / 21+ by average position."),
        ("AI Overview likely", "Position <= 5 and CTR below 40% of the expected CTR for that position: the answer is "
                               "probably being summarised above the results. A signal, not proof."),
        ("Est. extra clicks", "Impressions x expected CTR at the target position (1 for top 3, 3 for 4-10, 5 for 11-20, "
                              "8 for 21+) minus current clicks. Conservative curve: 25% / 15% / 10% / 7% / 5.5% for 1-5."),
        ("What GSC can't tell you", "Search Console doesn't label AI Mode, AI Overview or assistant traffic. These "
                                    "queries are found by their shape (length, phrasing, intent), so treat the list as "
                                    "the best available proxy. Very rare queries are anonymised by Google and never "
                                    "appear in the table."),
        ("1,000-row cap", "The GSC interface shows up to 1,000 rows per table. Run single regex families, or the API, "
                          "for the full long tail."),
        ("AI Feature Pages", "From GSC > Performance > Generative AI features (Beta): impressions where a page appeared "
                             "in AI Overviews / AI Mode. By page, not by query. Page rows add up to more than the "
                             "report total because one AI answer can show several of your pages."),
        ("Follow-ups & agents", "Follow-ups: short prompts with no topic word ('show me photos', 'which is cheaper?') typed "
                                "inside an AI conversation. Agent searches: site:/intitle:/quoted queries from AI tools; "
                                "reported, not optimised for."),
    ]
    for w in d.get("warnings", []):
        notes.append(("Data warning", w))
    for i, (k, v) in enumerate(notes, 4):
        ws_me[f"A{i}"], ws_me[f"B{i}"] = k, v
        ws_me[f"A{i}"].font = Font(bold=True, size=10, color=INK)
        ws_me[f"B{i}"].alignment = WRAP

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = {"Summary": CYAN, "Optimization Plan": GREEN, "AI Visibility Plan": GREEN,
                                        "Method & Notes": MUTED}.get(ws.title, VIOLET)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"Saved {args.out}  ({len(queries)} queries, {len(clusters)} topics, {len(opp)} opportunities)")
    if unwritten:
        print(f"WARN: {unwritten} High/Medium topic(s) have no plan in strategy.json")

    if args.md:
        lines = [f"# AI Queries Finder - {site_name}", "", f"*{period}*", ""]
        ai_feat = tot.get("ai_feature_total") or tot.get("ai_feature_impressions")
        if ai_feat:
            lines.append(f"- **{int(ai_feat):,} impressions in Google's AI features** on "
                         f"{tot.get('ai_feature_pages') or 0} pages (Generative AI features report)")
        lines += [f"- **{tot['ai_queries']:,} AI-style queries** ({tot['ai_impressions']:,} impressions, "
                  f"{tot['ai_clicks']:,} clicks), incl. {tot.get('followup_queries', 0)} AI Mode follow-ups"]
        if site_impr:
            lines.append(f"- **{tot['ai_impressions'] / site_impr:.1%}** of the site's impressions")
        lines += [f"- **{tot['aio_likely']}** queries where an AI answer likely takes the click",
                  f"- **~{tot['est_extra_clicks']:,} extra clicks** per period if the plan works", ""]
        if st.get("headline"):
            lines += [st["headline"], ""]
        lines += ["## Top topics", "", "| Topic | Priority | Impressions | Avg position | Move |", "|---|---|---:|---:|---|"]
        for c in [c for c in clusters if c["head"] != "(misc)"][:5]:
            lines.append(f"| {c['label']} | {c['priority']} | {c['impressions']:,} | {c['position']} | "
                         f"{cplan.get(c['id'], {}).get('page_action', '')} |")
        if wins:
            lines += ["", "## Quick wins (days 1-7)", ""] + [f"- {w.get('action', '')}" for w in wins]
        lines += ["", f"*{ATTRIBUTION}*", ""]
        args.md.parent.mkdir(parents=True, exist_ok=True)
        args.md.write_text("\n".join(lines), encoding="utf-8")
        print(f"Saved {args.md}")


if __name__ == "__main__":
    main()

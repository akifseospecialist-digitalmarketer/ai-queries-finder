#!/usr/bin/env python3
"""
verify_workbook.py - Step 8 of the ai-queries-finder skill.

Recalculates the workbook with LibreOffice (if installed) and checks:
  * every expected tab exists
  * no formula errors (#REF!, #VALUE!, #DIV/0!, #NAME?, #N/A)
  * every regex in the Regex Library is GSC-safe (RE2: no lookarounds/backreferences) and <= 4,096 chars
  * answer-first blocks are 40-60 words
  * no "(not written yet)" plans left for High/Medium topics (warning only)

Usage:  python verify_workbook.py outputs/report.xlsx
Exit code 1 if a hard check fails.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from find_ai_queries import check_regex  # noqa: E402

TABS = ["Summary", "AI Queries", "Regex Library", "Topic Clusters", "Opportunities", "Optimization Plan",
        "AI Visibility Plan", "Method & Notes"]
ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "Err:")
NOT_WRITTEN = "(not written yet)"


def recalc(path: Path) -> tuple[Path | None, str]:
    """Returns (recalculated copy, status) where status is 'ok', 'missing' or 'failed'."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                  "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
            if Path(p).exists():
                soffice = p
                break
    if not soffice:
        return None, "missing"
    out = Path(tempfile.mkdtemp())
    profile = Path(tempfile.mkdtemp())  # private profile: parallel runs can't fight over a lock
    try:
        res = subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--calc",
                              "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                             capture_output=True, timeout=240)
    except subprocess.TimeoutExpired:
        return None, "failed"
    f = out / path.name
    return (f, "ok") if res.returncode == 0 and f.exists() else (None, "failed")


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = Path(sys.argv[1])
    hard, soft = [], []

    wb = load_workbook(path)
    for t in TABS:
        if t not in wb.sheetnames:
            hard.append(f"Missing tab: {t}")

    if "Regex Library" in wb.sheetnames:
        ws = wb["Regex Library"]
        n = 0
        for r in range(5, ws.max_row + 1):
            rx = ws[f"C{r}"].value
            if not rx:
                continue
            n += 1
            for p in check_regex(str(rx)):
                hard.append(f"Regex {ws[f'A{r}'].value}: {p}")
        if not n:
            hard.append("Regex Library is empty")

    if "Optimization Plan" in wb.sheetnames:
        ws = wb["Optimization Plan"]
        empty = 0
        for r in range(2, ws.max_row + 1):
            block = ws[f"G{r}"].value
            if block == NOT_WRITTEN:
                empty += 1
            elif block:
                words = len(str(block).split())
                if not 40 <= words <= 60:
                    hard.append(f"Answer block for {ws[f'A{r}'].value} is {words} words (want 40-60)")
        if empty:
            soft.append(f"{empty} High/Medium topic plan(s) not written yet")

    if "AI Queries" in wb.sheetnames and wb["AI Queries"].max_row < 2:
        soft.append("AI Queries tab is empty - no queries matched")

    calc, status = recalc(path)
    if status == "failed":
        hard.append("LibreOffice is installed but recalculation failed - formulas were NOT checked. Re-run verify.")
    if calc:
        wbv = load_workbook(calc, data_only=True)
        for sh in wbv.worksheets:
            for row in sh.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith(ERRORS):
                        hard.append(f"Formula error {c.value} at {sh.title}!{c.coordinate}")
        s = wbv["Summary"]
        print("Summary KPIs after recalculation:")
        for r in range(4, 14):
            if s[f"A{r}"].value:
                v = s[f"B{r}"].value
                v = f"{v:.2%}" if isinstance(v, float) and v < 1 and "CTR" in str(s[f"A{r}"].value) + \
                    ("Share" if "Share" in str(s[f"A{r}"].value) else "") else v
                print(f"  {s[f'A{r}'].value:<52} {v}")
    elif status == "missing":
        soft.append("LibreOffice not found - formulas not recalculated (they will calculate when opened in Excel/Sheets)")

    for m in soft:
        print("WARN :", m)
    for m in hard:
        print("FAIL :", m)
    print("RESULT:", "FAIL" if hard else "PASS")
    sys.exit(1 if hard else 0)


if __name__ == "__main__":
    main()
